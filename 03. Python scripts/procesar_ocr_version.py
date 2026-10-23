# -*- coding: utf-8 -*-
import os
import sys
import time
import threading
import traceback
import shutil
import subprocess
from datetime import datetime
import pymupdf as fitz
from openpyxl import Workbook, load_workbook

# ==========================================================
# ID UNICO DE PROCESO (YYYYMMDDHHMMSS)
# Se genera al iniciar y se graba en todos los registros
# (05. Logs\ocr_log_<ID>.txt, ocr_log_<ID>.xlsx
# y 04. Resultados\Estadisticas OCR\ocr_estadisticas.xlsx)
# ==========================================================

ID_PROCESO = datetime.now().strftime("%Y%m%d%H%M%S")

# ==========================================================
# CONFIGURACION
# ==========================================================

DIRECTORIO_SCRIPT = os.path.dirname(os.path.abspath(__file__))
DIRECTORIO_RAIZ = os.path.dirname(DIRECTORIO_SCRIPT)

# Carpeta real de facturas: vive en la raiz de la biblioteca "Facturacion Inteligente (IA)",
# al lado de "98. AUTOMATIZACION" (DIRECTORIO_RAIZ), no dentro de ella.
NOMBRE_CARPETA_FACTURAS = "01. FACTURAS PENDIENTES CONTABILIZAR (AI)"
CARPETA_FACTURAS = os.path.join(
    os.path.dirname(DIRECTORIO_RAIZ), NOMBRE_CARPETA_FACTURAS
)
CARPETA_RESULTADOS = os.path.join(DIRECTORIO_RAIZ, "04. Resultados")
CARPETA_ESTADISTICAS = os.path.join(CARPETA_RESULTADOS, "Estadisticas OCR")

# Todos los logs del proceso OCR se guardan directamente en 05. Logs
CARPETA_LOGS = os.path.join(DIRECTORIO_RAIZ, "05. Logs")

# No se crea una subcarpeta por corrida: el ID ya forma parte del nombre
# de cada archivo de log.
CARPETA_CORRIDA = CARPETA_LOGS

# Carpeta temporal donde quedan los PDFs con OCR de ESTA corrida (se replica la
# estructura de carpetas) hasta que los devuelve devolucion_ocr_facturas_version.py
PREFIJO_OCR_SALIDA = "OCR_PARA_DEVOLVER"
CARPETA_OCR_SALIDA = os.path.join(CARPETA_LOGS, f"{PREFIJO_OCR_SALIDA}_{ID_PROCESO}")

os.makedirs(CARPETA_RESULTADOS, exist_ok=True)
os.makedirs(CARPETA_ESTADISTICAS, exist_ok=True)
os.makedirs(CARPETA_LOGS, exist_ok=True)
# (la carpeta OCR_PARA_DEVOLVER_<ID> se crea solo si hay al menos un OCR)

# Un log nuevo por cada corrida, identificado con el ID de proceso
ARCHIVO_LOG = os.path.join(CARPETA_CORRIDA, f"ocr_log_{ID_PROCESO}.xlsx")
ARCHIVO_LOG_TXT = os.path.join(CARPETA_CORRIDA, f"ocr_log_{ID_PROCESO}.txt")
ARCHIVO_ESTADISTICAS = os.path.join(CARPETA_ESTADISTICAS, "ocr_estadisticas.xlsx")
# # Maestro historico de fechas (acumulativo, fuente de verdad: solo se le agregan
# # registros) y su copia de respaldo, que se pisa en cada corrida.
# CARPETA_FECHAS = os.path.join(CARPETA_RESULTADOS, "Maestro de Fechas")
# ARCHIVO_MAESTRO_FECHAS = os.path.join(CARPETA_FECHAS, "maestro_fechas_facturas.xlsx")
# ARCHIVO_MAESTRO_FECHAS_COPIA = os.path.join(CARPETA_FECHAS, "maestro_fechas_facturas_copia.xlsx")
# # Ubicacion de la version anterior (se migra una sola vez si existe)
# ARCHIVO_MAESTRO_FECHAS_ANTERIOR = os.path.join(CARPETA_RESULTADOS, "maestro_fechas_facturas.xlsx")

# Maestro historico de fechas (acumulativo, fuente de verdad: solo se le agregan registros).
CARPETA_MAESTROS = os.path.join(DIRECTORIO_RAIZ, "02. Archivos Maestros")
CARPETA_FECHAS = os.path.join(CARPETA_MAESTROS, "Maestro Unificado")
ARCHIVO_MAESTRO_FECHAS = os.path.join(CARPETA_FECHAS, "MAESTRO_FECHAS_FACTURAS.xlsx")
# Ubicacion anterior (04. Resultados\Maestro de Fechas): se migra una sola vez
# si el maestro todavia no existe en la ubicacion nueva
ARCHIVO_MAESTRO_FECHAS_ANTERIOR = os.path.join(
    CARPETA_RESULTADOS, "Maestro de Fechas", "maestro_fechas_facturas.xlsx"
)
JOBS = 4

# ==========================================================
# TESSERACT DEL PROYECTO
# ==========================================================

# Carpeta de dependencias: se busca en este orden (la primera que tenga tesseract.exe).
# La actual es "01. Dependencias"; "Dependencias" queda como respaldo por si la
# carpeta vieja sigue en otra PC.
NOMBRES_CARPETA_DEPENDENCIAS = ("01. Dependencias", "Dependencias")


def buscar_tesseract():
    """Devuelve (carpeta_tesseract, [rutas_revisadas])."""
    revisadas = []
    for nombre in NOMBRES_CARPETA_DEPENDENCIAS:
        carpeta = os.path.join(DIRECTORIO_RAIZ, nombre, "Tesseract")
        revisadas.append(carpeta)
        if os.path.isfile(os.path.join(carpeta, "tesseract.exe")):
            return carpeta, revisadas
    return None, revisadas


TESSERACT_DIR, _TESSERACT_REVISADAS = buscar_tesseract()

if TESSERACT_DIR is None:
    print("ERROR: no se encontro Tesseract (tesseract.exe). Se busco en:", flush=True)
    for _ruta in _TESSERACT_REVISADAS:
        print("  - " + _ruta, flush=True)
    print("Verifique que OneDrive termino de sincronizar la carpeta "
          "'01. Dependencias' (Mantener siempre en este dispositivo).", flush=True)
    sys.exit(1)

TESSDATA_DIR = os.path.join(TESSERACT_DIR, "tessdata")

os.environ["PATH"] += os.pathsep + TESSERACT_DIR
os.environ["TESSDATA_PREFIX"] = TESSDATA_DIR

# ==========================================================
# MENSAJES EN PANTALLA + LOG DE TEXTO EN VIVO
# ==========================================================

def formato_tiempo(segundos):
    segundos = int(segundos)
    minutos, seg = divmod(segundos, 60)
    horas, minutos = divmod(minutos, 60)
    if horas:
        return f"{horas}:{minutos:02d}:{seg:02d}"
    return f"{minutos:02d}:{seg:02d}"


def formato_hhmmss(segundos):
    segundos = int(segundos)
    minutos, seg = divmod(segundos, 60)
    horas, minutos = divmod(minutos, 60)
    return f"{horas:02d}:{minutos:02d}:{seg:02d}"


def mensaje(texto=""):
    """Muestra el mensaje en pantalla y lo agrega al log de texto."""
    print(texto, flush=True)
    try:
        with open(ARCHIVO_LOG_TXT, "a", encoding="utf-8") as f:
            hora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            f.write(f"[{hora}] [{ID_PROCESO}] {texto}\n")
    except Exception:
        pass

# ==========================================================
# LOG XLSX
# ==========================================================

ENCABEZADO = [
    "Fecha",
    "Archivo",
    "Estado",
    "Detalle",
    "DuracionSegundos",
    "RutaOrigen",   # PDF original (de donde se extrajo)
    "RutaActual",   # Donde queda el archivo (PDF con OCR) hasta ser devuelto
    "IdProceso"     # YYYYMMDDHHMMSS de la ejecucion
]


def guardar_log(filas):
    if not filas:
        return

    if os.path.isfile(ARCHIVO_LOG):
        wb = load_workbook(ARCHIVO_LOG)
        ws = wb.active

        # Migracion: si el log viejo no tiene las columnas nuevas, se agregan
        for i, titulo in enumerate(ENCABEZADO, start=1):
            if ws.cell(row=1, column=i).value != titulo:
                ws.cell(row=1, column=i).value = titulo
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Log"
        ws.append(ENCABEZADO)

    for fila in filas:
        ws.append(list(fila) + [ID_PROCESO])

    wb.save(ARCHIVO_LOG)


def volcar_log(filas_log):
    """Guarda las filas pendientes en el Excel. Si el Excel esta abierto,
    las conserva en memoria y reintenta mas adelante."""
    try:
        guardar_log(filas_log)
        filas_log.clear()
    except PermissionError:
        mensaje(f"AVISO: no se pudo escribir {os.path.basename(ARCHIVO_LOG)} (esta abierto en Excel?). "
                "Cierrelo para que se guarde el registro.")
    except Exception as e:
        mensaje(f"AVISO: no se pudo guardar {os.path.basename(ARCHIVO_LOG)}: {e}")

# ==========================================================
# ESTADISTICAS XLSX (modo append, en 04. Resultados\Estadisticas OCR)
# ==========================================================

ENCABEZADO_ESTADISTICAS = [
    "Fecha",
    "Hora Inicio",
    "Hora Fin",
    "Duracion (hh:mm:ss)",
    "Duracion (min)",
    "Estado Proceso",
    "PDFs Encontrados",
    "Con Texto",
    "Sin Texto",
    "OCR Realizados",
    "OCR Existentes",
    "Errores",
    "IdProceso"
]


def guardar_estadisticas(fila):
    """Agrega una fila al historial de estadisticas (no sobreescribe)."""
    if os.path.isfile(ARCHIVO_ESTADISTICAS):
        wb = load_workbook(ARCHIVO_ESTADISTICAS)
        ws = wb.active

        # Migracion: si faltan columnas nuevas en el encabezado, se agregan
        for i, titulo in enumerate(ENCABEZADO_ESTADISTICAS, start=1):
            if ws.cell(row=1, column=i).value != titulo:
                ws.cell(row=1, column=i).value = titulo
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Estadisticas"
        ws.append(ENCABEZADO_ESTADISTICAS)

    ws.append(list(fila) + [ID_PROCESO])
    wb.save(ARCHIVO_ESTADISTICAS)


def volcar_estadisticas(fila):
    try:
        guardar_estadisticas(fila)
    except PermissionError:
        # Respaldo para no perder el registro si el Excel esta abierto
        respaldo = os.path.join(
            CARPETA_ESTADISTICAS,
            f"ocr_estadisticas_PENDIENTE_{ID_PROCESO}.xlsx"
        )
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Estadisticas"
            ws.append(ENCABEZADO_ESTADISTICAS)
            ws.append(list(fila) + [ID_PROCESO])
            wb.save(respaldo)
            mensaje("AVISO: ocr_estadisticas.xlsx esta abierto en Excel. "
                    f"El registro se guardo en: {respaldo}")
        except Exception as e:
            mensaje(f"AVISO: no se pudo guardar las estadisticas: {e}")
    except Exception as e:
        mensaje(f"AVISO: no se pudo guardar las estadisticas: {e}")

# ==========================================================
# MAESTRO HISTORICO DE FECHAS (acumulativo, nunca se sobrescribe)
# Clave de control: ARCHIVO + RUTA_COMPLETA
# ==========================================================

ENCABEZADO_MAESTRO_FECHAS = [
    "ARCHIVO",
    "RUTA_COMPLETA",
    "FECHA_RECEPCION",              # creacion del archivo en Windows
    "FECHA_MODIFICACION_ARCHIVO",   # ultima modificacion en Windows
    "FECHA_ULTIMO_ACCESO",          # ultimo acceso en Windows
    "FECHA_PROCESAMIENTO"           # momento en que este script detecto la factura
]

FORMATO_FECHA_EXCEL = "yyyy-mm-dd hh:mm:ss"


def clave_maestro(archivo, ruta):
    """Clave ARCHIVO + RUTA_COMPLETA (sin distinguir mayusculas, como Windows)."""
    return (
        os.path.normcase(str(archivo).strip()),
        os.path.normcase(os.path.normpath(str(ruta).strip()))
    )


def fechas_windows(ruta):
    """Devuelve (creacion, modificacion, ultimo acceso) del archivo."""
    st = os.stat(ruta)
    # Python 3.12+ en Windows: st_birthtime. Versiones anteriores: st_ctime
    # es la fecha de creacion en Windows.
    creacion = getattr(st, "st_birthtime", None)
    if creacion is None:
        creacion = st.st_ctime
    return (
        datetime.fromtimestamp(creacion),
        datetime.fromtimestamp(st.st_mtime),
        datetime.fromtimestamp(st.st_atime)
    )


def _formatear_hoja_maestro(ws):
    ws.freeze_panes = "A2"
    for letra, ancho in zip("ABCDEF", (45, 90, 22, 28, 22, 22)):
        ws.column_dimensions[letra].width = ancho


def _agregar_fila_maestro(ws, fila):
    ws.append(fila)
    for col in (3, 4, 5, 6):
        ws.cell(row=ws.max_row, column=col).number_format = FORMATO_FECHA_EXCEL


def _guardar_atomico(wb, destino):
    """Guarda en un temporal y luego reemplaza, para no dejar archivos a medio escribir."""
    temporal = destino + ".tmp"
    try:
        wb.save(temporal)
        os.replace(temporal, destino)
    except Exception:
        borrar_si_existe(temporal)
        raise


def _abrir_maestro():
    """Devuelve (wb, ws, claves). Crea el maestro en memoria si no existe.
    Si existe la version anterior (04. Resultados, con ID_PROCESO) se migran
    sus registros sin la columna ID_PROCESO; el archivo viejo no se toca."""
    claves = set()

    if os.path.isfile(ARCHIVO_MAESTRO_FECHAS):
        wb = load_workbook(ARCHIVO_MAESTRO_FECHAS)
        ws = wb.active

        encabezado = [ws.cell(row=1, column=i).value
                      for i in range(1, len(ENCABEZADO_MAESTRO_FECHAS) + 1)]
        if encabezado != ENCABEZADO_MAESTRO_FECHAS:
            # Fuente de verdad: si el formato no es el esperado no se toca.
            raise ValueError(
                "el encabezado del maestro no coincide con el esperado; "
                "no se modifico el archivo"
            )

        for fila in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if fila[0] is None and fila[1] is None:
                continue
            claves.add(clave_maestro(fila[0] or "", fila[1] or ""))
        return wb, ws, claves

    wb = Workbook()
    ws = wb.active
    ws.title = "Maestro"
    ws.append(ENCABEZADO_MAESTRO_FECHAS)
    _formatear_hoja_maestro(ws)

    if os.path.isfile(ARCHIVO_MAESTRO_FECHAS_ANTERIOR):
        wb_ant = load_workbook(ARCHIVO_MAESTRO_FECHAS_ANTERIOR, read_only=True)
        n = len(ENCABEZADO_MAESTRO_FECHAS)
        for fila in wb_ant.active.iter_rows(min_row=2, max_col=n, values_only=True):
            if fila[0] is None and fila[1] is None:
                continue
            clave = clave_maestro(fila[0] or "", fila[1] or "")
            if clave in claves:
                continue
            _agregar_fila_maestro(ws, list(fila))
            claves.add(clave)
        wb_ant.close()

    return wb, ws, claves


# def registrar_en_maestro_fechas(rutas_pdf):
#     """Agrega al maestro SOLO las facturas cuya clave ARCHIVO + RUTA_COMPLETA
#     no existe todavia. Conserva todos los registros historicos.
#     Devuelve (nuevos, ya_existentes, sin_fechas, copia_ok)."""
#     os.makedirs(CARPETA_FECHAS, exist_ok=True)

#     existia = os.path.isfile(ARCHIVO_MAESTRO_FECHAS)
#     wb, ws, claves = _abrir_maestro()
#     migrados = 0 if existia else len(claves)

#     nuevos = 0
#     existentes = 0
#     sin_fechas = 0

#     for ruta in rutas_pdf:
#         archivo = os.path.basename(ruta)
#         clave = clave_maestro(archivo, ruta)

#         if clave in claves:
#             existentes += 1
#             continue

#         try:
#             recepcion, modificacion, acceso = fechas_windows(ruta)
#         except OSError:
#             sin_fechas += 1
#             continue

#         _agregar_fila_maestro(ws, [
#             archivo, ruta, recepcion, modificacion, acceso, datetime.now()
#         ])
#         claves.add(clave)
#         nuevos += 1

#     if nuevos or migrados:
#         _guardar_atomico(wb, ARCHIVO_MAESTRO_FECHAS)

#     # Copia de respaldo: queda igual al maestro tal como termino esta corrida.
#     # Solo se llega aca si el maestro se leyo y valido bien; si fallo antes, la
#     # copia anterior queda intacta.
#     copia_ok = True
#     if os.path.isfile(ARCHIVO_MAESTRO_FECHAS):
#         try:
#             temporal = ARCHIVO_MAESTRO_FECHAS_COPIA + ".tmp"
#             try:
#                 shutil.copy2(ARCHIVO_MAESTRO_FECHAS, temporal)
#                 os.replace(temporal, ARCHIVO_MAESTRO_FECHAS_COPIA)
#             except Exception:
#                 borrar_si_existe(temporal)
#                 raise
#         except Exception:
#             copia_ok = False

#     return nuevos, existentes, sin_fechas, copia_ok

def registrar_en_maestro_fechas(rutas_pdf):
    """Agrega al maestro SOLO las facturas cuya clave ARCHIVO + RUTA_COMPLETA
    no existe todavia. Conserva todos los registros historicos.
    Devuelve (nuevos, ya_existentes, sin_fechas)."""
    os.makedirs(CARPETA_FECHAS, exist_ok=True)
    existia = os.path.isfile(ARCHIVO_MAESTRO_FECHAS)
    wb, ws, claves = _abrir_maestro()
    migrados = 0 if existia else len(claves)
    nuevos = 0
    existentes = 0
    sin_fechas = 0
    for ruta in rutas_pdf:
        archivo = os.path.basename(ruta)
        clave = clave_maestro(archivo, ruta)
        if clave in claves:
            existentes += 1
            continue
        try:
            recepcion, modificacion, acceso = fechas_windows(ruta)
        except OSError:
            sin_fechas += 1
            continue
        _agregar_fila_maestro(ws, [
            archivo, ruta, recepcion, modificacion, acceso, datetime.now()
        ])
        claves.add(clave)
        nuevos += 1
    if nuevos or migrados:
        _guardar_atomico(wb, ARCHIVO_MAESTRO_FECHAS)
    return nuevos, existentes, sin_fechas

def actualizar_maestro_fechas(rutas_pdf):
    nombre = os.path.basename(ARCHIVO_MAESTRO_FECHAS)
    try:
        nuevos, existentes, sin_fechas = registrar_en_maestro_fechas(rutas_pdf)
    except PermissionError:
        mensaje(f"AVISO: no se pudo actualizar {nombre} (esta abierto en Excel?). "
                "Cierrelo y vuelva a ejecutar: las facturas pendientes "
                "se registraran en la proxima corrida.")
        return
    except Exception as e:
        mensaje(f"AVISO: no se actualizo {nombre}: {e}")
        return
    mensaje(f"Maestro de fechas: {nuevos} facturas nuevas registradas, "
            f"{existentes} ya estaban registradas.")
    if sin_fechas:
        mensaje(f"AVISO: no se pudieron leer las fechas de {sin_fechas} archivos.")

# ==========================================================
# TEXTO
# ==========================================================

def tiene_texto(pdf):
    try:
        doc = fitz.open(pdf)

        for pagina in doc:
            texto = pagina.get_text()

            if texto and texto.strip():
                doc.close()
                return True

        doc.close()

    except Exception:
        pass

    return False

# ==========================================================
# OCR (con contador de tiempo en vivo)
# ==========================================================

def ejecutar_ocr(pdf_entrada, pdf_salida):
    comando = [
        sys.executable,
        "-m",
        "ocrmypdf",
        "--skip-text",
        "--rotate-pages",
        "--jobs",
        str(JOBS),
        "-l",
        "spa",
        pdf_entrada,
        pdf_salida
    ]

    proceso = subprocess.Popen(
        comando,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=os.environ.copy()
    )

    salida = []

    def leer():
        for linea in proceso.stdout:
            salida.append(linea)

    hilo = threading.Thread(target=leer, daemon=True)
    hilo.start()

    inicio = time.time()

    try:
        while proceso.poll() is None:
            print(
                f"\r      Trabajando... {formato_tiempo(time.time() - inicio)}   ",
                end="",
                flush=True
            )
            time.sleep(1)
    except KeyboardInterrupt:
        proceso.kill()
        raise
    finally:
        print("\r" + " " * 50 + "\r", end="", flush=True)

    hilo.join(timeout=5)

    return proceso.returncode, "".join(salida)


def indexar_ocr_previos():
    """Devuelve {ruta relativa del OCR: ruta completa} con los OCR que ya
    estan en carpetas OCR_PARA_DEVOLVER* de corridas anteriores (backup)."""
    indice = {}

    for nombre in sorted(os.listdir(CARPETA_LOGS)):
        ruta = os.path.join(CARPETA_LOGS, nombre)

        if not os.path.isdir(ruta):
            continue
        if not nombre.startswith(PREFIJO_OCR_SALIDA):
            continue
        if os.path.normcase(ruta) == os.path.normcase(CARPETA_OCR_SALIDA):
            continue

        for raiz, _, archivos in os.walk(ruta):
            for a in archivos:
                if a.lower().endswith("_ocr.pdf"):
                    completa = os.path.join(raiz, a)
                    clave = os.path.normcase(os.path.relpath(completa, ruta))
                    indice.setdefault(clave, completa)

    return indice


def borrar_si_existe(ruta):
    try:
        if os.path.isfile(ruta):
            os.remove(ruta)
    except Exception:
        pass

# ==========================================================
# MAIN
# ==========================================================

def main():
    inicio_global = time.time()
    inicio_dt = datetime.now()

    filas_log = []

    encontrados = 0
    con_texto = 0
    sin_texto = 0
    ocr_realizados = 0
    ocr_existentes = 0
    errores = 0
    tiempo_ocr_total = 0.0
    interrumpido = False

    mensaje("====================================")
    mensaje("OCR DE FACTURAS")
    mensaje(f"ID de proceso: {ID_PROCESO}")
    mensaje(f"Carpeta de logs de esta corrida: {CARPETA_CORRIDA}")
    mensaje("====================================")
    if not os.path.isdir(CARPETA_FACTURAS):
        mensaje("ERROR: no se encuentra la carpeta de facturas:")
        mensaje(CARPETA_FACTURAS)
        mensaje("Verifique que este sincronizada y que el nombre sea exacto.")
        return 1
    mensaje("Buscando PDFs en la carpeta de facturas...")

    # ------------------------------------------------------
    # Listado previo de PDFs para poder mostrar el avance
    # ------------------------------------------------------
    pendientes = []

    for raiz, _, archivos in os.walk(CARPETA_FACTURAS):
        for archivo in sorted(archivos):
            nombre = archivo.lower()

            if not nombre.endswith(".pdf"):
                continue

            if nombre.endswith("_ocr.pdf"):
                continue

            ruta_pdf = os.path.join(raiz, archivo)
            nombre_base = os.path.splitext(archivo)[0]
            relativo = os.path.relpath(ruta_pdf, CARPETA_FACTURAS)
            carpeta_relativa = os.path.dirname(relativo)

            # OCR ya devuelto (junto al original)
            ruta_ocr_original = os.path.join(raiz, nombre_base + "_OCR.pdf")

            # OCR generado en la carpeta de salida (pendiente de devolver)
            ruta_ocr_salida = os.path.join(
                CARPETA_OCR_SALIDA, carpeta_relativa, nombre_base + "_OCR.pdf"
            )

            pendientes.append(
                (ruta_pdf, ruta_ocr_original, ruta_ocr_salida, relativo)
            )

    encontrados = len(pendientes)
    ocr_previos = indexar_ocr_previos()

    mensaje(f"Se encontraron {encontrados} PDFs para revisar.")
    mensaje(f"Los PDFs con OCR se guardan en: {CARPETA_OCR_SALIDA}")
    mensaje("(Para detener el proceso en forma segura: Ctrl + C)")
    mensaje("------------------------------------")

    # Maestro historico de fechas: se registra toda factura detectada antes
    # de empezar el OCR, asi queda guardada aunque el proceso se interrumpa.
    mensaje("Actualizando maestro de fechas de facturas...")
    actualizar_maestro_fechas([p[0] for p in pendientes])
    mensaje("------------------------------------")

    try:
        for numero, (ruta_pdf, ruta_ocr_original, ruta_ocr_salida, relativo) in enumerate(pendientes, start=1):
            prefijo = f"[{numero}/{encontrados}]"

            # OCR ya devuelto a la carpeta original
            if os.path.isfile(ruta_ocr_original):
                ocr_existentes += 1
                mensaje(f"{prefijo} {relativo} -> ya tenia OCR, se omite")
                filas_log.append([
                    datetime.now(), relativo, "OCR_YA_EXISTE", "", 0,
                    ruta_pdf, ruta_ocr_original
                ])
                continue

            # OCR ya generado en una corrida anterior (carpeta de backup)
            clave = os.path.normcase(os.path.splitext(relativo)[0] + "_OCR.pdf")
            ruta_previa = ocr_previos.get(clave)

            if ruta_previa:
                ocr_existentes += 1
                corrida_previa = os.path.relpath(
                    ruta_previa, CARPETA_LOGS
                ).split(os.sep)[0]
                mensaje(f"{prefijo} {relativo} -> OCR ya generado en {corrida_previa}, se omite")
                filas_log.append([
                    datetime.now(), relativo, "OCR_EN_CORRIDA_ANTERIOR",
                    corrida_previa, 0, ruta_pdf, ruta_previa
                ])
                continue

            try:
                if tiene_texto(ruta_pdf):
                    con_texto += 1
                    mensaje(f"{prefijo} {relativo} -> ya tiene texto, se omite")
                    filas_log.append([
                        datetime.now(), relativo, "CON_TEXTO", "", 0,
                        ruta_pdf, ruta_pdf
                    ])
                    continue

                sin_texto += 1

                mensaje(f"{prefijo} {relativo} -> aplicando OCR...")

                os.makedirs(os.path.dirname(ruta_ocr_salida), exist_ok=True)

                inicio = time.time()

                try:
                    codigo, salida = ejecutar_ocr(ruta_pdf, ruta_ocr_salida)
                except KeyboardInterrupt:
                    borrar_si_existe(ruta_ocr_salida)
                    raise

                duracion = round(time.time() - inicio, 2)

                if codigo == 0:
                    ocr_realizados += 1
                    tiempo_ocr_total += duracion

                    promedio = tiempo_ocr_total / ocr_realizados

                    mensaje(
                        f"      LISTO en {formato_tiempo(duracion)} "
                        f"(promedio por PDF: {formato_tiempo(promedio)})"
                    )

                    filas_log.append([
                        datetime.now(), relativo, "OCR_REALIZADO", "", duracion,
                        ruta_pdf, ruta_ocr_salida
                    ])
                else:
                    errores += 1

                    # Evita dejar un archivo _OCR.pdf incompleto
                    borrar_si_existe(ruta_ocr_salida)

                    detalle = salida[-500:]

                    mensaje(
                        f"      ERROR al procesar este PDF "
                        f"(ver detalle en {os.path.basename(ARCHIVO_LOG)} / {os.path.basename(ARCHIVO_LOG_TXT)})"
                    )
                    mensaje(f"      Detalle: {detalle.strip()}")

                    filas_log.append([
                        datetime.now(), relativo, "ERROR", detalle, duracion,
                        ruta_pdf, ruta_pdf
                    ])

                volcar_log(filas_log)

            except KeyboardInterrupt:
                raise

            except Exception as e:
                errores += 1
                mensaje(f"      ERROR inesperado en este PDF: {e}")
                filas_log.append([
                    datetime.now(), relativo, "ERROR", str(e), 0,
                    ruta_pdf, ruta_pdf
                ])

    except KeyboardInterrupt:
        interrumpido = True
        mensaje("")
        mensaje("*** Proceso detenido por el usuario. Se guarda lo realizado hasta ahora. ***")

    # ------------------------------------------------------
    # Guardado final
    # ------------------------------------------------------
    volcar_log(filas_log)

    fin_dt = datetime.now()
    duracion_total = time.time() - inicio_global

    # Estadisticas en modo append (04. Resultados\Estadisticas OCR)
    volcar_estadisticas([
        inicio_dt.strftime("%Y-%m-%d"),
        inicio_dt.strftime("%H:%M:%S"),
        fin_dt.strftime("%H:%M:%S"),
        formato_hhmmss(duracion_total),
        round(duracion_total / 60, 2),
        "INTERRUMPIDO" if interrumpido else "COMPLETO",
        encontrados,
        con_texto,
        sin_texto,
        ocr_realizados,
        ocr_existentes,
        errores
    ])

    tiempo_total = round(duracion_total / 60, 2)

    mensaje("")
    mensaje("====================================")
    mensaje("RESUMEN OCR" + (" (INTERRUMPIDO)" if interrumpido else ""))
    mensaje("====================================")
    mensaje(f"PDFs encontrados: {encontrados}")
    mensaje(f"Con texto: {con_texto}")
    mensaje(f"Sin texto: {sin_texto}")
    mensaje(f"OCR realizados: {ocr_realizados}")
    mensaje(f"OCR existentes: {ocr_existentes}")
    mensaje(f"Errores: {errores}")
    mensaje(f"Tiempo total: {tiempo_total} min")
    mensaje("====================================")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception:
        mensaje("")
        mensaje("*** Ocurrio un error inesperado. Avise al area de sistemas con este detalle: ***")
        mensaje(traceback.format_exc())
        raise SystemExit(1)
